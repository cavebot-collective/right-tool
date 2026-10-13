from pathlib import Path
from zipfile import ZipFile, ZIP_STORED, ZIP_DEFLATED
from docx import Document

ROOT = Path(__file__).resolve().parent
FIX = ROOT / 'fixtures'
FIX.mkdir(exist_ok=True)
MARKER = 'RIGHT_TOOL_PANDOC_MARKER_314159'
TITLE = 'Right Tool Bakeoff'
P1 = f'This paragraph contains {MARKER} for extraction verification.'
ITEMS = ['Alpha item', 'Beta item']

(FIX/'sample.md').write_text(f'# {TITLE}\n\n{P1}\n\n- {ITEMS[0]}\n- {ITEMS[1]}\n', encoding='utf-8')
(FIX/'sample.html').write_text(f'<!doctype html><html><body><h1>{TITLE}</h1><p>{P1}</p><ul><li>{ITEMS[0]}</li><li>{ITEMS[1]}</li></ul></body></html>', encoding='utf-8')
(FIX/'sample.tex').write_text(r'''\section{''' + TITLE + r'''}

''' + P1 + r'''

\begin{itemize}
\item ''' + ITEMS[0] + r'''
\item ''' + ITEMS[1] + r'''
\end{itemize}
''', encoding='utf-8')
(FIX/'sample.rst').write_text(f'{TITLE}\n{"="*len(TITLE)}\n\n{P1}\n\n- {ITEMS[0]}\n- {ITEMS[1]}\n', encoding='utf-8')

doc = Document()
doc.add_heading(TITLE, level=1)
doc.add_paragraph(P1)
for item in ITEMS:
    doc.add_paragraph(item, style='List Bullet')
doc.save(FIX/'sample.docx')

# python-docx writes current ZIP member timestamps. Canonicalize them so the
# generated DOCX is byte-reproducible while preserving every member's content.
from zipfile import ZipInfo
docx = FIX/'sample.docx'
tmp = FIX/'sample.docx.tmp'
with ZipFile(docx, 'r') as src, ZipFile(tmp, 'w') as dst:
    for info in src.infolist():
        data = src.read(info.filename)
        z = ZipInfo(info.filename, (2020,1,1,0,0,0))
        z.compress_type = info.compress_type
        z.external_attr = info.external_attr
        z.internal_attr = info.internal_attr
        z.create_system = info.create_system
        z.flag_bits = info.flag_bits
        dst.writestr(z, data)
tmp.replace(docx)

# Minimal deterministic EPUB. ZIP timestamps are fixed for byte reproducibility.
FIXED=(2020,1,1,0,0,0)
def zi(name, compress=ZIP_DEFLATED):
    from zipfile import ZipInfo
    z=ZipInfo(name, FIXED); z.compress_type=compress; return z
with ZipFile(FIX/'sample.epub','w') as z:
    z.writestr(zi('mimetype', ZIP_STORED), 'application/epub+zip')
    z.writestr(zi('META-INF/container.xml'), '''<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>''')
    z.writestr(zi('OEBPS/content.xhtml'), f'''<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>{TITLE}</title></head><body><h1>{TITLE}</h1><p>{P1}</p><ul><li>{ITEMS[0]}</li><li>{ITEMS[1]}</li></ul></body></html>''')
    z.writestr(zi('OEBPS/content.opf'), '''<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" unique-identifier="bookid" version="3.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="bookid">urn:right-tool:bakeoff</dc:identifier><dc:title>Right Tool Bakeoff</dc:title><dc:language>en</dc:language></metadata><manifest><item id="content" href="content.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="content"/></spine></package>''')

print(MARKER)
for p in sorted(FIX.iterdir()): print(p.name, p.stat().st_size)
