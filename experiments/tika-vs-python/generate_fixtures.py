from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
FIX = ROOT / "fixtures"
FIX.mkdir(exist_ok=True)

MARKER = "RIGHT_TOOL_TIKA_MARKER_271828"
TITLE = "Right Tool Tika Bakeoff"
TEXT = f"{TITLE} — {MARKER}"

(FIX / "sample.txt").write_text(TEXT + "\n", encoding="utf-8")
(FIX / "sample.html").write_text(
    f"<!doctype html><html><body><h1>{TITLE}</h1><p>{MARKER}</p></body></html>",
    encoding="utf-8",
)
(FIX / "sample.xml").write_text(
    f'<?xml version="1.0" encoding="utf-8"?><root><title>{TITLE}</title><marker>{MARKER}</marker></root>',
    encoding="utf-8",
)
(FIX / "sample.rtf").write_text(
    r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Courier;}}\f0\fs24 "
    + TITLE
    + r"\par "
    + MARKER
    + r"\par}",
    encoding="ascii",
)

FIXED = (2020, 1, 1, 0, 0, 0)

def zi(name: str, compression=ZIP_DEFLATED) -> ZipInfo:
    info = ZipInfo(name, FIXED)
    info.compress_type = compression
    return info

with ZipFile(FIX / "sample.zip", "w") as z:
    z.writestr(zi("inside.txt"), TEXT + "\n")

doc = Document()
doc.add_heading(TITLE, level=1)
doc.add_paragraph(MARKER)
doc.save(FIX / "sample.docx")

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws["A1"] = TITLE
ws["A2"] = MARKER
wb.save(FIX / "sample.xlsx")

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
box.text_frame.text = TITLE + "\n" + MARKER
prs.save(FIX / "sample.pptx")

pdf = canvas.Canvas(str(FIX / "sample.pdf"))
pdf.drawString(72, 760, TITLE)
pdf.drawString(72, 730, MARKER)
pdf.save()

with ZipFile(FIX / "sample.epub", "w") as z:
    z.writestr(zi("mimetype", ZIP_STORED), "application/epub+zip")
    z.writestr(
        zi("META-INF/container.xml"),
        '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',
    )
    z.writestr(
        zi("OEBPS/content.xhtml"),
        f'<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>{TITLE}</title></head><body><h1>{TITLE}</h1><p>{MARKER}</p></body></html>',
    )
    z.writestr(
        zi("OEBPS/content.opf"),
        '<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" unique-identifier="bookid" version="3.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="bookid">urn:right-tool:tika-bakeoff</dc:identifier><dc:title>Right Tool Tika Bakeoff</dc:title><dc:language>en</dc:language></metadata><manifest><item id="content" href="content.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="content"/></spine></package>',
    )

for path in sorted(FIX.iterdir()):
    print(path.name, path.stat().st_size)
