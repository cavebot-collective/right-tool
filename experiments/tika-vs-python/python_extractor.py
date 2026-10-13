from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
from striprtf.striprtf import rtf_to_text

def extract(path: Path) -> str:
    ext = path.suffix.lower()

    if ext == ".txt":
        return path.read_text(encoding="utf-8")

    if ext == ".html":
        return BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser").get_text("\n")

    if ext == ".xml":
        root = ET.fromstring(path.read_text(encoding="utf-8"))
        return "\n".join(text for text in root.itertext())

    if ext == ".rtf":
        return rtf_to_text(path.read_text(encoding="ascii"))

    if ext == ".zip":
        with ZipFile(path) as z:
            return "\n".join(
                z.read(name).decode("utf-8", "replace")
                for name in z.namelist()
                if name.lower().endswith(".txt")
            )

    if ext == ".docx":
        return "\n".join(p.text for p in Document(path).paragraphs)

    if ext == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        return "\n".join(
            str(value)
            for ws in wb.worksheets
            for row in ws.iter_rows(values_only=True)
            for value in row
            if value is not None
        )

    if ext == ".pptx":
        prs = Presentation(path)
        return "\n".join(
            shape.text
            for slide in prs.slides
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text
        )

    if ext == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)

    if ext == ".epub":
        with ZipFile(path) as z:
            docs = [
                name for name in z.namelist()
                if name.lower().endswith((".xhtml", ".html", ".htm"))
            ]
            return "\n".join(
                BeautifulSoup(z.read(name), "html.parser").get_text("\n")
                for name in docs
            )

    raise ValueError(f"unsupported format: {ext}")
